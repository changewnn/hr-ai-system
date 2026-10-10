from io import BytesIO

import pytest
from docx import Document

from app.services.knowledge import ParsedSegment, chunk_segments, parse_document, prepare_upload


def test_prepare_upload_accepts_supported_text_file() -> None:
    prepared = prepare_upload("policy.md", "text/markdown", b"# Leave policy\nEmployees may apply.")

    assert prepared.extension == ".md"
    assert len(prepared.checksum) == 64


def test_prepare_upload_rejects_executable_extension() -> None:
    with pytest.raises(ValueError, match="UNSUPPORTED_TYPE"):
        prepare_upload("policy.exe", "application/octet-stream", b"MZ")


def test_prepare_upload_rejects_mismatched_pdf_signature() -> None:
    with pytest.raises(ValueError, match="INVALID_SIGNATURE"):
        prepare_upload("policy.pdf", "application/pdf", b"not a pdf")


def test_markdown_parser_keeps_section_titles(tmp_path) -> None:
    path = tmp_path / "policy.md"
    path.write_text(
        "# Annual Leave\nTen days per year.\n## Request\nApply in advance.", encoding="utf-8"
    )

    segments = parse_document(str(path))

    assert [segment.section for segment in segments] == ["Annual Leave", "Request"]
    assert "Ten days" in segments[0].text


def test_docx_parser_keeps_heading(tmp_path) -> None:
    document = Document()
    document.add_heading("Attendance", level=1)
    document.add_paragraph("Employees check in before 09:00.")
    stream = BytesIO()
    document.save(stream)
    path = tmp_path / "policy.docx"
    path.write_bytes(stream.getvalue())

    segments = parse_document(str(path))

    assert segments[0].section == "Attendance"
    assert "09:00" in segments[0].text


def test_chunking_preserves_metadata_and_overlap() -> None:
    source = ParsedSegment(text="A" * 1200, section="Section A", page_number=2)

    chunks = chunk_segments([source], chunk_size=500, overlap=100)

    assert len(chunks) == 3
    assert all(chunk.section == "Section A" and chunk.page_number == 2 for chunk in chunks)
    assert chunks[0].text[-100:] == chunks[1].text[:100]
